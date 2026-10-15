"""格式分类与转换引擎（doc-compile）。

设计要点：
- 不信任扩展名，按「魔数」判断真实文件类型（很多 .wps 实际是 OOXML zip）。
- 分类结果：
    zip     → OOXML（.docx/.docm，及 WPS 以 OOXML 保存的 .wps），python-docx 直接读
    ole2    → OLE2 复合文档（老版 .doc 等），经 LibreOffice 转 docx
    rtf     → RTF，经 LibreOffice 转 docx
    text    → 纯文本，python-docx 组装
    unknown → 其它二进制（含真私有 .wps），尝试 LibreOffice，失败则丢弃
- PDF 相关（预览渲染、页码定位）优先使用 poppler（pdftotext/pdftoppm/pdfinfo），
  docx→pdf 由 LibreOffice 负责。sugon（mips64el）无 PyMuPDF wheel，故不依赖 fitz。
"""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document

from ... import config

# ---- 魔数 ----
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
RTF_MAGIC = b"{\\rtf"

_soffice_cache: Optional[str] = None
_poppler_re = re.compile(r"^Pages:\s+(\d+)", re.M)


def sniff(data: bytes) -> str:
    """按文件头判定真实类型：zip | ole2 | rtf | text | unknown。"""
    if len(data) < 4:
        return "text" if data.strip() else "unknown"
    if data[:2] == b"PK":
        return "zip"
    if data[:8] == OLE2_MAGIC:
        return "ole2"
    if data[:5].lower() == RTF_MAGIC:
        return "rtf"
    # 常见 UTF-8/GBK/UTF-16 文本启发式：不含大量控制字符
    sample = data[:4096]
    if b"\x00\x00" in sample[:16]:
        return "unknown"
    try:
        sample.decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        try:
            sample.decode("gbk")
            return "text"
        except UnicodeDecodeError:
            return "unknown"


def classify(data: bytes) -> str:
    """返回对外分类名：docx | doc | rtf | plain | unknown。"""
    kind = sniff(data)
    if kind == "zip":
        return "docx"  # 含 .wps 的 OOXML 变体
    if kind == "ole2":
        return "doc"
    if kind == "rtf":
        return "rtf"
    if kind == "text":
        return "plain"
    return "unknown"


# ---------------------------------------------------------------------------
# LibreOffice 探测
# ---------------------------------------------------------------------------

def find_soffice() -> Optional[str]:
    global _soffice_cache
    if _soffice_cache is not None:
        return _soffice_cache or None
    candidates: List[str] = []
    if config.SOFFICE_PATH:
        candidates.append(config.SOFFICE_PATH)
    for name in ("soffice", "libreoffice", "lowriter"):
        found = shutil.which(name)
        if found:
            candidates.append(found)
    win_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    candidates.extend(p for p in win_paths if os.path.exists(p))
    for cand in candidates:
        if Path(cand).is_file():
            _soffice_cache = cand
            return cand
    _soffice_cache = ""
    return None


def soffice_available() -> bool:
    return find_soffice() is not None


def soffice_version() -> str:
    soffice = find_soffice()
    if not soffice:
        return ""
    try:
        proc = subprocess.run(
            [soffice, "--version"], capture_output=True, text=True, timeout=15,
        )
        return (proc.stdout or proc.stderr).strip().splitlines()[0][:80]
    except Exception:
        return ""


def convert_with_soffice(src_bytes: bytes, src_suffix: str, out_format: str,
                         timeout: Optional[int] = None) -> Optional[bytes]:
    """用 LibreOffice headless 将 src 转换（docx/pdf/doc/…），返回产物字节或 None。"""
    soffice = find_soffice()
    if not soffice:
        return None
    timeout = timeout or config.COMPILE_CONVERT_TIMEOUT
    with tempfile.TemporaryDirectory(prefix="dc_cvt_") as tmp:
        src_path = Path(tmp) / ("input" + src_suffix)
        src_path.write_bytes(src_bytes)
        out_dir = Path(tmp) / "out"
        out_dir.mkdir()
        try:
            proc = subprocess.run(
                [soffice, "--headless", "--norestore", "--convert-to", out_format,
                 "--outdir", str(out_dir), str(src_path)],
                capture_output=True, timeout=timeout,
            )
        except subprocess.TimeoutExpired:
            return None
        if proc.returncode != 0:
            return None
        for candidate in out_dir.iterdir():
            if candidate.is_file():
                return candidate.read_bytes()
        return None


# ---------------------------------------------------------------------------
# docx 相关
# ---------------------------------------------------------------------------

def valid_docx(data: bytes) -> bool:
    """校验 bytes 是否为可被 python-docx 打开的 OOXML。"""
    try:
        Document(io.BytesIO(data))
        return True
    except Exception:
        return False


def rewrite_first_paragraphs(docx_bytes: bytes, title: str, author: str) -> bytes:
    """改写前两个非空段落的文本（保留原段落样式），返回新的 docx bytes。"""
    doc = Document(io.BytesIO(docx_bytes))
    replaced = 0
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        new_text = title if replaced == 0 else author
        # 清空原有 runs，写入单 run（保守方案：仅当内容确实需要替换）
        if para.text.strip() != new_text:
            for run in list(para.runs):
                run._r.getparent().remove(run._r)
            para.add_run(new_text)
        replaced += 1
        if replaced >= 2:
            break
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def text_to_docx_bytes(text: str) -> bytes:
    """把纯文本组装为 docx：首行=标题（Heading 1），次行=作者，其余为正文。"""
    doc = Document()
    lines = [ln.strip() for ln in text.replace("\r\n", "\n").split("\n")]
    lines = [ln for ln in lines if ln]
    if lines:
        doc.add_heading(lines[0], level=1)
    if len(lines) > 1:
        p = doc.add_paragraph(lines[1])
        p.alignment = 1  # CENTER
    for ln in lines[2:]:
        doc.add_paragraph(ln)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# poppler（PDF 文本/渲染，不依赖 PyMuPDF）
# ---------------------------------------------------------------------------

def _bin(name: str) -> Optional[str]:
    return shutil.which(name)


def poppler_available() -> bool:
    return bool(_bin("pdftotext") and _bin("pdftoppm") and _bin("pdfinfo"))


def poppler_version() -> str:
    try:
        proc = subprocess.run(
            [_bin("pdftotext") or "pdftotext", "-v"],
            capture_output=True, text=True, timeout=10,
        )
        return (proc.stderr or proc.stdout).strip().splitlines()[0][:80]
    except Exception:
        return ""


def pdfinfo_pages(pdf_bytes: bytes) -> Optional[int]:
    pdfinfo_bin = _bin("pdfinfo")
    if not pdfinfo_bin:
        return None
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(pdf_bytes)
        tmp = f.name
    try:
        proc = subprocess.run([pdfinfo_bin, tmp], capture_output=True, text=True, timeout=30)
        m = _poppler_re.search(proc.stdout or "")
        return int(m.group(1)) if m else None
    except Exception:
        return None
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass


def pdftoppm_page_png(pdf_bytes: bytes, page: int, dpi: int) -> Optional[bytes]:
    pdftoppm = _bin("pdftoppm")
    if not pdftoppm:
        return None
    with tempfile.TemporaryDirectory(prefix="dc_png_") as tmp:
        pdf_path = Path(tmp) / "in.pdf"
        pdf_path.write_bytes(pdf_bytes)
        prefix = Path(tmp) / "page"
        try:
            proc = subprocess.run(
                [pdftoppm, "-f", str(page), "-l", str(page), "-r", str(dpi), "-png",
                 str(pdf_path), str(prefix)],
                capture_output=True, timeout=60,
            )
            if proc.returncode != 0:
                return None
        except subprocess.TimeoutExpired:
            return None
        matches = sorted(Path(tmp).glob("page*.png"))
        if not matches:
            return None
        return matches[0].read_bytes()


def pdftotext_page(pdf_bytes: bytes, page: int) -> str:
    pdftotext = _bin("pdftotext")
    if not pdftotext:
        return ""
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(pdf_bytes)
        tmp = f.name
    try:
        proc = subprocess.run(
            [pdftotext, "-f", str(page), "-l", str(page), "-layout", tmp, "-"],
            capture_output=True, timeout=30,
        )
        if proc.returncode != 0:
            return ""
        return proc.stdout.decode("utf-8", errors="ignore")
    except Exception:
        return ""
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass


def normalize_text(text: str) -> str:
    """归一化文本用于页内匹配（去空白）。"""
    return re.sub(r"\s+", "", text or "")


def locate_page_nums(total_pages: int, pdf_bytes: bytes, needles: List[str]) -> List[Optional[int]]:
    """按顺序在 PDF 各页中定位每个标题的起始页（归一化匹配，容忍页眉）。

    返回与 needles 等长的列表，未命中的位置为 None。
    """
    pages = [normalize_text(pdftotext_page(pdf_bytes, p)) for p in range(1, total_pages + 1)]
    results: List[Optional[int]] = []
    start = 0
    for needle in needles:
        key = normalize_text(needle)
        if not key or len(key) > 20:
            key = key[:20]
        found = None
        for i in range(start, len(pages)):
            if key and key in pages[i]:
                found = i + 1
                start = i
                break
        results.append(found)
    return results