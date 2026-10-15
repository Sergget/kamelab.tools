"""目录模板与静态目录生成（doc-compile）。

内置模板：
- <seq>. <title>……<author>              （序号. 标题……作者）
- <seq>. <title>【<author>】……<page>     （序号. 标题【作者】……页码）

「……」等引导符在 Word 中实现为制表位点线（dot leader），因此模板被拆成
「左侧文本 + 右侧文本」两段：左侧含 <seq>/<title>/<author>，右侧含
<author>/<page>。
"""

from __future__ import annotations

import re
from typing import List, Optional, Tuple

from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

DEFAULT_TEMPLATES = [
    {"label": "序号. 标题……作者", "value": "<seq>. <title>……<author>"},
    {"label": "序号. 标题【作者】……页码", "value": "<seq>. <title>【<author>】……<page>"},
]

# 引导符：连续的“……”/… / . - ─ 等
_LEADER_RE = re.compile(r"(?:…+|\.{3,}|-{3,}|—{3,}|─{3,}|\u3000+)", re.S)


def split_template(template: str) -> Tuple[str, str, bool]:
    """拆分模板为 (左侧模板, 右侧模板, 是否含页码)。

    引导符起分割作用；未找到引导符时整段为左侧，右侧为空。
    """
    tpl = (template or "").strip()
    if not tpl:
        tpl = "<seq>. <title>……<author>"
    m = _LEADER_RE.search(tpl)
    if not m:
        return tpl, "", "<page>" in tpl
    left = tpl[: m.start()].strip()
    right = tpl[m.end():].strip()
    with_page = "<page>" in right or (not right and "<page>" in tpl)
    return left, right, with_page


_EMPTY_BRACKET_RE = re.compile(r"[【（(〔]\s*[】）)〕]")


def _clean(s: str) -> str:
    s = _EMPTY_BRACKET_RE.sub("", s)
    return re.sub(r"<[^>]+>", "", s).strip()


def render_left(left_tpl: str, seq: int, title: str, author: str = "") -> str:
    out = (left_tpl.replace("<seq>", str(seq))
                  .replace("<title>", title or "")
                  .replace("<author>", author or "")
                  .replace("<page>", ""))
    return _clean(out)


def render_right(right_tpl: str, author: str, page: Optional[int]) -> str:
    out = (right_tpl.replace("<author>", author or "")
                    .replace("<page>", str(page) if page is not None else "")
                    .replace("<seq>", "").replace("<title>", ""))
    return _clean(out)


def format_toc_line(template: str, seq: int, title: str, author: str,
                    page: Optional[int]) -> Tuple[str, str]:
    """按模板渲染一行目录：返回 (左侧文本, 右侧文本)。"""
    left_tpl, right_tpl, with_page = split_template(template)
    left = render_left(left_tpl, seq, title, author)
    if with_page and page is None:
        right_tpl = right_tpl.replace("<page>", "").replace("【<author>】", "")
    right = render_right(right_tpl, author, page)
    # 无引导符但要求页码的模板：把页码直接放到右侧兜底
    if with_page and not right and page is not None:
        right = str(page)
    return left, right


def add_custom_toc_entry(doc, left: str, right: str, tab_cm: float = 15.5,
                         style: Optional[str] = None) -> None:
    """向 doc 追加一行带点线引导的目录（左侧 + 制表位 + 右侧）。"""
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(_cm(tab_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r1 = p.add_run(left)
    if right:
        p.add_run("\t" + right)


def _cm(value: float):
    from docx.shared import Cm
    return Cm(value)


def build_custom_toc(doc, entries: List[dict], template: str, tab_cm: float = 15.5) -> None:
    """entries: [{"seq","title","author","page"}] 按顺序渲染到 doc。"""
    for entry in entries:
        left, right = format_toc_line(
            template, entry["seq"], entry.get("title", ""),
            entry.get("author", ""), entry.get("page"),
        )
        add_custom_toc_entry(doc, left, right, tab_cm=tab_cm)