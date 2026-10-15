"""文档汇编工具测试（doc-compile）。

依赖：python-docx（必须）；LibreOffice/poppler 依赖的用例（预览图片、页码、
.doc 输出）在缺少相应工具时跳过。
"""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.tools.doc_compile import converters, detect, merge, service, toc

client = TestClient(app)


# ---------------------------------------------------------------------------
# 工具注册
# ---------------------------------------------------------------------------

def test_tools_registry_contains_doc_compile():
    r = client.get("/api/tools")
    ids = [t["id"] for t in r.json()]
    assert "doc-compile" in ids


# ---------------------------------------------------------------------------
# 魔数分类
# ---------------------------------------------------------------------------

def test_sniff_magic_bytes():
    assert converters.sniff(b"PK\x03\x04something") == "zip"
    assert converters.sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1more") == "ole2"
    assert converters.sniff(b"{\\rtf1\\ansi hello") == "rtf"
    assert converters.sniff("你好世界\n正文".encode("utf-8")) == "text"
    assert converters.sniff(b"") == "unknown"


def test_classify_mapping():
    assert converters.classify(b"PK\x03\x04x") == "docx"
    assert converters.classify(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") == "doc"
    assert converters.classify(b"{\\rtf1") == "rtf"
    assert converters.classify("纯文本内容".encode("utf-8")) == "plain"
    assert converters.classify(b"\x00\x01\x02\xff") == "unknown"


# ---------------------------------------------------------------------------
# 标题 / 作者识别
# ---------------------------------------------------------------------------

def test_detect_basic():
    title, author, auto, warns = detect.detect_title_author(
        ["关于人工智能的思考", "张三", "正文第一段", "正文第二段"]
    )
    assert title == "关于人工智能的思考"
    assert author == "张三"
    assert auto is True


def test_detect_skips_meta_line():
    title, author, _, _ = detect.detect_title_author(
        ["2024年5月", "关于人工智能的思考", "张三"]
    )
    assert title == "关于人工智能的思考"
    assert author == "张三"


def test_detect_author_prefix_stripped():
    _, author, _, _ = detect.detect_title_author(["标题", "作者：李四"])
    assert author == "李四"


def test_detect_author_second_line_is_body():
    title, author, _, warns = detect.detect_title_author(
        ["标题", "这是一段很长的正文内容，讲述了相关技术背景，并且还有更多细节描述，远远超过作者的合理长度。"]
    )
    assert author == ""
    assert any("作者" in w for w in warns)


def test_detect_title_from_filename_fallback():
    title, _, _, warns = detect.detect_title_author([], "调研报告.docx")
    assert title == "调研报告"
    assert any("文件名" in w for w in warns)


# ---------------------------------------------------------------------------
# 目录模板
# ---------------------------------------------------------------------------

def test_toc_split_default_templates():
    assert toc.split_template("<seq>. <title>……<author>") == (
        "<seq>. <title>", "<author>", False,
    )
    left, right, with_page = toc.split_template("<seq>. <title>【<author>】……<page>")
    assert with_page is True
    assert "【<author>】" in left
    assert right == "<page>"


def test_toc_render_two_templates():
    left, right = toc.format_toc_line("<seq>. <title>……<author>", 3, "关于XX的报告", "张三", None)
    assert left == "3. 关于XX的报告"
    assert right == "张三"

    left, right = toc.format_toc_line(
        "<seq>. <title>【<author>】……<page>", 5, "论云计算", "李四", 12,
    )
    assert left == "5. 论云计算【李四】"
    assert right == "12"


def test_toc_render_empty_author_brackets_cleaned():
    left, right = toc.format_toc_line(
        "<seq>. <title>【<author>】……<page>", 6, "无人标题", "", 3,
    )
    assert "【】" not in left
    assert left == "6. 无人标题"
    assert right == "3"


def test_toc_render_unknown_tokens_stripped():
    left, right = toc.format_toc_line("编号 <seq>｜<title>｜<author>｜页<page>", 1, "A文", "B人", 7)
    assert "<" not in left and ">" not in left
    assert "A文" in left and "B人" in left


# ---------------------------------------------------------------------------
# 合并（本地，不依赖 LibreOffice）
# ---------------------------------------------------------------------------

def make_docx(title: str, author: str, body: list) -> bytes:
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(author)
    for line in body:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_merge_two_docs_no_duplicate():
    from app.tools.doc_compile.models import CompileOptions, TocOptions

    items = [
        {"title": "第一篇", "author": "张三", "docx_bytes": make_docx("第一篇", "张三", ["正文甲。"]),
         "skip_texts": ["第一篇", "张三"]},
        {"title": "第二篇", "author": "李四", "docx_bytes": make_docx("第二篇", "李四", ["正文乙。"]),
         "skip_texts": ["第二篇", "李四"]},
    ]
    opts = CompileOptions(page_break_between=True)
    doc = merge.build_compiled_docx(items, opts)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts.count("第一篇") == 1
    assert texts.count("第二篇") == 1
    assert "正文甲。" in texts and "正文乙。" in texts


def test_merge_cover_and_custom_toc():
    from app.tools.doc_compile.models import CompileOptions, CoverOptions, TocOptions

    items = [
        {"title": "篇一", "author": "甲", "docx_bytes": make_docx("篇一", "甲", ["内容。"]),
         "skip_texts": ["篇一", "甲"]},
    ]
    opts = CompileOptions(
        cover=CoverOptions(enabled=True, title="汇编集", date="2026-01-01"),
        toc=TocOptions(enabled=True, mode="custom",
                       template="<seq>. <title>……<author>", include_page_numbers=False),
    )
    doc = merge.build_compiled_docx(items, opts)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts[0] == "汇编集"
    assert any("1. 篇一" in t for t in texts)


def test_merge_toc_field_mode():
    from app.tools.doc_compile.models import CompileOptions, TocOptions

    items = [
        {"title": "篇一", "author": "甲", "docx_bytes": make_docx("篇一", "甲", ["内容。"]),
         "skip_texts": ["篇一", "甲"]},
    ]
    opts = CompileOptions(toc=TocOptions(enabled=True, mode="field"))
    doc = merge.build_compiled_docx(items, opts)
    xml = doc.element.body.xml
    assert "instrText" in xml and "TOC" in xml
    assert "updateFields" in doc.settings.element.xml


def test_merge_keeps_table():
    from app.tools.doc_compile.models import CompileOptions

    d = Document()
    d.add_heading("表格篇", 0)
    d.add_paragraph("作者")
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "X"
    buf = io.BytesIO()
    d.save(buf)
    items = [
        {"title": "表格篇", "author": "作者", "docx_bytes": buf.getvalue(),
         "skip_texts": ["表格篇", "作者"]},
    ]
    doc = merge.build_compiled_docx(items, CompileOptions())
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "X"


def test_merge_strong_skip_user_renamed_title():
    """用户把标题改成与原文完全不同时，原文自带标题/作者不再残留正文。"""
    from app.tools.doc_compile.models import CompileOptions

    raw = make_docx("旧标题", "旧作者", ["正文第一段。", "正文第二段。"])
    items = [{
        "title": "全新标题", "author": "新作者", "docx_bytes": raw,
        "skip_texts": ["旧标题", "旧作者", "全新标题", "新作者"],
        "strong_skip": True,
    }]
    doc = merge.build_compiled_docx(items, CompileOptions())
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts == ["全新标题", "新作者", "正文第一段。", "正文第二段。"]


def test_merge_strong_skip_author_only():
    """只改作者时原文旧作者段同样被消费。"""
    from app.tools.doc_compile.models import CompileOptions

    raw = make_docx("标题", "旧作者", ["正文。"])
    items = [{
        "title": "标题", "author": "新作者", "docx_bytes": raw,
        "skip_texts": ["标题", "旧作者", "标题", "新作者"],
        "strong_skip": True,
    }]
    doc = merge.build_compiled_docx(items, CompileOptions())
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts == ["标题", "新作者", "正文。"]


def test_merge_no_strong_skip_when_untitled():
    """原作者识别自文件名兜底（原文实无标题）时，不得误吞正文首段。"""
    from app.tools.doc_compile.models import CompileOptions

    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("这篇文档没有标题，首段就是正文并且内容较长，讲的是测试相关背景。")
    d.add_paragraph("第二段。")
    d.save(buf)
    items = [{
        "title": "文件名兜底", "author": "", "docx_bytes": buf.getvalue(),
        "skip_texts": ["文件名兜底", ""],
        "strong_skip": False,  # 服务端依据 title_from_filename 决定
    }]
    doc = merge.build_compiled_docx(items, CompileOptions())
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert len(texts) == 3  # 文件名标题 + 两段正文，都不被吞
    assert "没有标题" in texts[1]


# ---------------------------------------------------------------------------
# 服务层：上传
# ---------------------------------------------------------------------------

def test_process_upload_valid_docx():
    item, drop = service.process_upload(make_docx("标题", "作者", ["正文"]), "测试文档.docx")
    assert drop is None
    assert item["title"] == "标题"
    assert item["author"] == "作者"
    assert item["classified_as"] == "docx"


def test_process_upload_plain_text():
    item, drop = service.process_upload("第一行标题\n第二行作者\n正文内容".encode("utf-8"), "笔记.txt")
    assert drop is None
    assert item["title"] == "第一行标题"
    assert item["author"] == "第二行作者"


def test_process_upload_unsupported_ext():
    item, drop = service.process_upload(b"hello", "报表.xlsx")
    assert item is None
    assert "扩展名" in drop


def test_process_upload_empty():
    item, drop = service.process_upload(b"", "空.docx")
    assert item is None
    assert "空" in drop


def test_process_upload_bad_zip():
    item, drop = service.process_upload(b"PK\x03\x04 not a real zip at all", "坏文档.docx")
    assert item is None
    assert drop is not None


# ---------------------------------------------------------------------------
# 服务层：写回源文档（原生 docx 不依赖 LibreOffice）
# ---------------------------------------------------------------------------

def test_update_meta_apply_to_source():
    item, _ = service.process_upload(make_docx("原标题", "原作者", ["正文。"]), "源文档.docx")
    fid = item["file_id"]
    updated = service.update_meta(fid, "新标题", "新作者", apply_to_source=True)
    assert updated["source_updated"] is True
    assert updated["title"] == "新标题"
    from app import temp_files
    raw = temp_files.read_work_file(fid, "updated.docx")
    doc = Document(io.BytesIO(raw))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert texts[0] == "新标题"
    assert texts[1] == "新作者"


# ---------------------------------------------------------------------------
# 服务层：汇编
# ---------------------------------------------------------------------------

def _upload_two(client_items=True):
    r1 = service.process_upload(make_docx("文章甲", "张三", ["甲正文。"]), "甲.docx")
    r2 = service.process_upload(make_docx("文章乙", "李四", ["乙正文。"]), "乙.docx")
    return r1[0], r2[0]


def test_compile_custom_toc_and_download():
    from app.tools.doc_compile.models import CompileItem, CompileOptions, CompileRequest, TocOptions

    a, b = _upload_two()
    req = CompileRequest(
        items=[CompileItem(file_id=a["file_id"]), CompileItem(file_id=b["file_id"])],
        options=CompileOptions(
            toc=TocOptions(enabled=True, mode="custom",
                           template="<seq>. <title>……<author>", include_page_numbers=False),
        ),
    )
    res = service.compile_documents(req)
    assert res["download_name"] == "汇编文档.docx"
    assert res["stats"]["doc_count"] == 2
    assert res["stats"]["toc_entries"] == 2

    man = service.temp_files.get_job(res["job_id"])
    out = Path(man["dir"]) / man["download_name"]
    assert zipfile.is_zipfile(str(out))
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "目  录" in texts
    assert any("1. 文章甲" in t for t in texts)
    assert texts.count("文章甲") == 1


def test_compile_wps_output():
    from app.tools.doc_compile.models import CompileItem, CompileOptions, CompileRequest

    a, _ = _upload_two()
    req = CompileRequest(
        items=[CompileItem(file_id=a["file_id"])],
        options=CompileOptions(output_format="wps"),
    )
    res = service.compile_documents(req)
    assert res["download_name"].endswith(".wps")
    assert any("wps" in w for w in res["stats"]["warnings"])


def test_compile_unknown_format_rejected():
    from app.tools.doc_compile.models import CompileItem, CompileOptions, CompileRequest

    a, _ = _upload_two()
    req = CompileRequest(
        items=[CompileItem(file_id=a["file_id"])],
        options=CompileOptions(output_format="pdf"),
    )
    with pytest.raises(ValueError):
        service.compile_documents(req)


def test_compile_missing_items_rejected():
    from app.tools.doc_compile.models import CompileRequest

    with pytest.raises(ValueError):
        service.compile_documents(CompileRequest(items=[]))


# ---------------------------------------------------------------------------
# API 集成
# ---------------------------------------------------------------------------

def api_make_docx_bytes() -> bytes:
    return make_docx("API标题", "API作者", ["API正文。"])


def test_api_upload_multi():
    files = [
        ("files", ("甲.docx", api_make_docx_bytes(),
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("笔记.txt", "标题\n作者\n正文".encode("utf-8"), "text/plain")),
        ("files", ("坏文件.docx", b"PK\x03\x04 oops", "application/octet-stream")),
    ]
    r = client.post("/api/tools/doc-compile/upload", files=files)
    assert r.status_code == 200
    data = r.json()
    assert len(data["items"]) == 2
    assert len(data["dropped"]) == 1
    assert data["dropped"][0]["original_name"] == "坏文件.docx"


def test_api_upload_rejects_bad_extension():
    files = [("files", ("a.xlsx", b"hello", "application/octet-stream"))]
    r = client.post("/api/tools/doc-compile/upload", files=files)
    data = r.json()
    assert data["items"] == []
    assert "扩展名" in data["dropped"][0]["reason"]


def test_api_info_and_preview_text():
    files = [("files", ("甲.docx", api_make_docx_bytes(), "application/octet-stream"))]
    file_id = client.post("/api/tools/doc-compile/upload", files=files).json()["items"][0]["file_id"]

    info = client.get(f"/api/tools/doc-compile/files/{file_id}/info")
    assert info.status_code == 200
    assert info.json()["title"] == "API标题"

    pt = client.get(f"/api/tools/doc-compile/files/{file_id}/preview-text")
    assert pt.status_code == 200
    assert "API标题" in pt.json()["text"]


def test_api_meta_update():
    files = [("files", ("甲.docx", api_make_docx_bytes(), "application/octet-stream"))]
    file_id = client.post("/api/tools/doc-compile/upload", files=files).json()["items"][0]["file_id"]

    r = client.put(f"/api/tools/doc-compile/files/{file_id}/meta",
                   json={"title": "改后标题", "author": "改后作者", "apply_to_source": True})
    assert r.status_code == 200
    body = r.json()
    assert body["title"] == "改后标题"
    assert body["source_updated"] is True

    dl = client.get(f"/api/tools/doc-compile/files/{file_id}/source-download")
    assert dl.status_code == 200
    assert dl.headers["content-disposition"].startswith("attachment")


def test_api_compile_and_download():
    files = [
        ("files", ("甲.docx", api_make_docx_bytes(), "application/octet-stream")),
        ("files", ("乙.docx", make_docx("乙标题", "乙作者", ["乙正文。"]), "application/octet-stream")),
    ]
    ups = client.post("/api/tools/doc-compile/upload", files=files).json()["items"]
    payload = {
        "items": [{"file_id": ups[0]["file_id"]}, {"file_id": ups[1]["file_id"]}],
        "options": {"output_format": "docx",
                    "toc": {"enabled": True, "mode": "custom",
                            "template": "<seq>. <title>【<author>】……<page>",
                            "include_page_numbers": False}},
    }
    r = client.post("/api/tools/doc-compile/compile", json=payload)
    assert r.status_code == 200
    data = r.json()
    assert data["stats"]["doc_count"] == 2

    dl = client.get(f"/api/tools/doc-compile/download/{data['job_id']}")
    assert dl.status_code == 200
    assert len(dl.content) > 0


def test_api_compile_empty_items_rejected():
    r = client.post("/api/tools/doc-compile/compile", json={"items": []})
    assert r.status_code == 400


def test_api_supported_formats_and_health():
    fmts = client.get("/api/tools/doc-compile/supported-formats")
    assert fmts.status_code == 200
    body = fmts.json()
    assert ".docx" in body["input_formats"]
    assert body["output_formats"] == ["doc", "docx", "wps"]
    assert len(body["toc_templates"]) == 2

    h = client.get("/api/tools/doc-compile/health")
    assert h.json()["status"] == "UP"


# ---------------------------------------------------------------------------
# 依赖 LibreOffice / poppler 的用例（缺工具时跳过）
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not converters.soffice_available(),
                    reason="缺少 LibreOffice")
def test_api_preview_image():
    files = [("files", ("甲.docx", api_make_docx_bytes(), "application/octet-stream"))]
    file_id = client.post("/api/tools/doc-compile/upload", files=files).json()["items"][0]["file_id"]
    r = client.get(f"/api/tools/doc-compile/files/{file_id}/preview?page=1")
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/png"


@pytest.mark.skipif(not (converters.soffice_available() and converters.poppler_available()),
                    reason="缺少 LibreOffice 或 poppler")
def test_compile_custom_toc_with_page_numbers():
    from app.tools.doc_compile.models import CompileItem, CompileOptions, CompileRequest, TocOptions

    a, b = _upload_two()
    req = CompileRequest(
        items=[CompileItem(file_id=a["file_id"]), CompileItem(file_id=b["file_id"])],
        options=CompileOptions(
            toc=TocOptions(enabled=True, mode="custom",
                           template="<seq>. <title>【<author>】……<page>",
                           include_page_numbers=True),
        ),
    )
    res = service.compile_documents(req)
    man = service.temp_files.get_job(res["job_id"])
    doc = Document(str(Path(man["dir"]) / man["download_name"]))
    toc_texts = [p.text for p in doc.paragraphs[:6] if p.text.strip()]
    assert any("\t" in t for t in toc_texts)


@pytest.mark.skipif(not converters.soffice_available(),
                    reason="缺少 LibreOffice")
def test_compile_doc_output():
    from app.tools.doc_compile.models import CompileItem, CompileOptions, CompileRequest

    a, _ = _upload_two()
    req = CompileRequest(
        items=[CompileItem(file_id=a["file_id"])],
        options=CompileOptions(output_format="doc"),
    )
    res = service.compile_documents(req)
    assert res["download_name"] == "汇编文档.doc"
    man = service.temp_files.get_job(res["job_id"])
    out = Path(man["dir"]) / man["download_name"]
    assert out.read_bytes()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # OLE2 头