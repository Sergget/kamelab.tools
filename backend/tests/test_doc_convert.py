"""文档转换工具测试。"""

import io
import zipfile

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.tools.doc_convert import factory, local_extract, service

client = TestClient(app)


# ---- crop 解析与校验 ----

def test_parse_crop_valid():
    crop = service.parse_crop('{"unit":"ratio","x":0.1,"y":0.2,"width":0.5,"height":0.3}')
    assert crop == {"x": 0.1, "y": 0.2, "width": 0.5, "height": 0.3}


def test_parse_crop_none():
    assert service.parse_crop("") is None
    assert service.parse_crop(None) is None


@pytest.mark.parametrize("raw", [
    "not-json",
    '{"unit":"pixel","x":10}',
    '{"unit":"ratio","x":1.5,"y":0,"width":0.1,"height":0.1}',
    '{"unit":"ratio","x":0.6,"y":0,"width":0.5,"height":0.1}',
    '{"unit":"ratio","x":0,"y":0,"width":0,"height":0.1}',
    '{"unit":"ratio","x":0,"y":0,"width":-1,"height":0.1}',
])
def test_parse_crop_invalid(raw):
    with pytest.raises(service.DocConvertError):
        service.parse_crop(raw)


# ---- 输出格式 ----

MARKDOWN_SAMPLE = (
    "# 大标题\n\n"
    "**加粗**与*斜体*，`行内代码`，[链接](http://example.com)，![图](img.png)。\n\n"
    "> 引用内容\n\n"
    "| 姓名 | 分数 |\n"
    "| --- | --- |\n"
    "| 张三 | 90 |\n\n"
    "---\n\n"
    "- 项目A\n"
    "- 项目B\n\n"
    "```python\n"
    "print('hi')\n"
    "```\n"
)


def test_strip_markdown_light():
    text = service.strip_markdown_light(MARKDOWN_SAMPLE)
    assert "# 大标题" not in text
    assert "大标题" in text
    assert "**" not in text and "*斜体*" not in text
    assert "行内代码" in text
    assert "[链接]" not in text and "链接" in text
    assert "> 引用内容" not in text and "引用内容" in text
    assert "| 张三 | 90 |" not in text
    assert "张三   90" in text
    assert "```" not in text


def test_markdown_to_docx_fallback():
    docx_bytes, engine = service.markdown_to_docx_bytes(MARKDOWN_SAMPLE)
    # 引擎取决于部署环境是否安装 pandoc，两种均合法
    assert engine in ("pandoc", "python-docx-fallback")
    assert len(docx_bytes) > 0
    assert zipfile.is_zipfile(io.BytesIO(docx_bytes))
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "大标题" in xml
    assert "张三" in xml
    assert "项目A" in xml


def test_validate_output_format():
    assert service.validate_output_format("txt") == "txt"
    with pytest.raises(service.DocConvertError):
        service.validate_output_format("html")


# ---- 版面重建（bbox 几何启发式）----

def make_line(text, top, height, left=10.0):
    return {
        "text": text,
        "bbox": [
            [left, top], [left + 200, top],
            [left + 200, top + height], [left, top + height],
        ],
    }


def test_reconstruct_layout_headings_and_paragraphs():
    median = 12.0
    lines = [
        make_line("一级大标题", 0, median * 2.5),           # 高度 ≥ 2x 中位 → #
        make_line("正文第一行", 40, median),
        make_line("正文第二行", 54, median),                 # 同行连续 → 合并
        make_line("二级标题", 80, median * 1.8),            # 1.6~2x → ##
        make_line("缩进内容", 110, median, left=40.0),      # 左缩进 → 引用块
        make_line("第三段", 140, median),
    ]
    md = service._reconstruct_layout(lines)
    assert md.startswith("# 一级大标题")
    assert "## 二级标题" in md
    assert "正文第一行正文第二行" in md  # 中文字符直接拼接不补空格
    assert "> 缩进内容" in md
    assert "\n\n" in md


def test_reconstruct_layout_paragraph_gap():
    lines = [
        make_line("第一段开头", 0, 12.0),
        make_line("第一段结尾", 15.0, 12.0),
        make_line("第二段开头", 60.0, 12.0),   # 间距 60-27=33 >> 12*1.4 → 分段
    ]
    md = service._reconstruct_layout(lines)
    assert "第一段开头第一段结尾" in md
    assert md.count("\n\n") == 1
    assert md.split("\n\n")[1] == "第二段开头"


# ---- PDF 分类与路由入口 ----

def test_classify_pdf_garbage_is_scanned():
    assert service.classify_pdf(b"not-a-real-pdf") == "scanned"


def test_convert_rejects_unknown_ext():
    with pytest.raises(service.DocConvertError):
        service.convert_document(b"hello", "文件.xyz", None, "md")


def test_convert_rejects_docx_to_docx():
    with pytest.raises(service.DocConvertError, match="docx"):
        service.convert_document(b"x", "源文档.docx", None, "docx")


def test_convert_rejects_bad_crop():
    with pytest.raises(service.DocConvertError):
        service.parse_crop("not-json")


# ---- 本地提取（纯 Python，MarkItDown 替代）----

def test_local_extract_md_txt_json():
    assert local_extract.extract_markdown("# hi\n\n正文".encode(), ".md").startswith("# hi")
    assert "内容" in local_extract.extract_markdown("内容".encode("gbk"), ".txt")
    assert '"a"' in local_extract.extract_markdown(b'{"a": 1}', ".json")


def test_local_extract_csv():
    md = local_extract.extract_markdown("姓名,分数\n张三,90\n李四,80\n".encode(), ".csv")
    assert "| 姓名 | 分数 |" in md
    assert "| 张三 | 90 |" in md
    assert "| 李四 | 80 |" in md


def test_local_extract_html():
    html = ("<html><head><style>.x{}</style></head><body>"
            "<h1>标题</h1><p>正文<b>加粗</b></p></body></html>").encode("utf-8")
    text = local_extract.extract_markdown(html, ".html")
    assert "标题" in text and "正文" in text and "<" not in text


def test_local_extract_docx():
    from docx import Document as MakeDocx

    doc = MakeDocx()
    doc.add_heading("测试标题", level=1)
    doc.add_paragraph("普通段落文字")
    buf = io.BytesIO()
    doc.save(buf)
    md = local_extract.extract_markdown(buf.getvalue(), ".docx")
    assert "测试标题" in md and "普通段落文字" in md


def test_local_extract_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["姓名", "分数"])
    ws.append(["张三", 90])
    buf = io.BytesIO()
    wb.save(buf)
    md = local_extract.extract_markdown(buf.getvalue(), ".xlsx")
    assert "## 成绩" in md and "张三" in md and "90" in md


def test_local_extract_pptx():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "幻灯片文字"
    buf = io.BytesIO()
    prs.save(buf)
    md = local_extract.extract_markdown(buf.getvalue(), ".pptx")
    assert "第 1 页" in md and "幻灯片文字" in md


def test_local_extract_unsupported():
    with pytest.raises(local_extract.LocalExtractError):
        local_extract.extract_markdown(b"x", ".rtf")


def test_convert_unsupported_office_ext_now_rejected():
    # 分支裁剪了 zip/epub/odt 等格式支持，应明确报错而不是静默失败
    r = client.post(
        "/api/tools/doc-convert/convert",
        files={"file": ("a.epub", b"xx", "application/epub+zip")},
    )
    assert r.status_code == 400


def test_convert_docx_endpoint_real_file():
    from docx import Document as MakeDocx

    doc = MakeDocx()
    doc.add_heading("接口测试", level=2)
    doc.add_paragraph("这是正文内容")
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post(
        "/api/tools/doc-convert/convert",
        files={"file": ("真实文档.docx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document")},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["success"] is True
    assert "接口测试" in (data["content"] or "")
    assert "这是正文内容" in (data["content"] or "")


# ---- HTTP 接口 ----

def md_bytes() -> bytes:
    return "# 标题\n\n正文内容\n".encode("utf-8")


def test_convert_md_endpoint_and_download():
    r = client.post(
        "/api/tools/doc-convert/convert",
        files={"file": ("示例.md", md_bytes(), "text/markdown")},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["success"] is True
    assert data["engine"] == "pure-python"
    assert data["content"].startswith("# 标题")
    assert data["download_name"] == "示例.md"

    dl = client.get(f"/api/tools/doc-convert/download/{data['job_id']}")
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith("text/markdown")
    assert "标题" in dl.content.decode("utf-8")


def test_convert_docx_endpoint_download():
    r = client.post(
        "/api/tools/doc-convert/convert",
        files={"file": ("示例.md", md_bytes(), "text/markdown")},
        data={"output_format": "docx"},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["docx_engine"] in ("pandoc", "python-docx-fallback")
    assert zipfile.is_zipfile(io.BytesIO(__import__("base64").b64decode(data["file_base64"])))

    dl = client.get(f"/api/tools/doc-convert/download/{data['job_id']}")
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith("application/vnd.openxmlformats-")


def test_convert_formats_endpoint():
    r = client.get("/api/tools/doc-convert/supported-formats")
    assert r.status_code == 200
    assert r.json()["output_formats"] == ["md", "txt", "docx"]
    assert ".docx" in r.json()["local_formats"]
    assert ".png" in r.json()["ocr_formats"]


def test_convert_health_endpoint():
    r = client.get("/api/tools/doc-convert/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "UP"
    assert "ocr_status" in body and "metrics" in body


def test_tools_registry_contains_doc_convert():
    tools = client.get("/api/tools").json()
    assert any(t["id"] == "doc-convert" and t["route"] == "/tools/doc-convert" for t in tools)


# ---- 回归测试：多页 OCR 结果聚合 ----

class _StubPdfExtractor(factory.BaseExtractor):
    """桩提取器：把每页渲染结果识别为固定的两行文字，用于验证多页聚合逻辑。"""

    def render_pdf_pages(self, file_bytes, dpi=200):
        return [(1, b"page-1"), (2, b"page-2"), (3, b"page-3")]

    def _ocr_image(self, img_bytes):
        tag = img_bytes.decode()
        return {
            "text": f"{tag}-A\n{tag}-B",
            "pages": 1,
            "lines": [
                {"text": f"{tag}-A", "bbox": [[0, 0], [10, 0], [10, 5], [0, 5]]},
                {"text": f"{tag}-B", "bbox": [[0, 6], [10, 6], [10, 11], [0, 11]]},
            ],
        }


def test_multi_page_ocr_aggregates_all_pages(monkeypatch):
    """多页 PDF 的 OCR 必须返回**所有页**的行与页码，而不是最后一页的行。

    历史 bug：循环结束后返回 result.get("lines") 与 len(all_texts)，
    只拿到最后一页的行，且 pages 统计的是「有文字页数」而非总页数。
    """
    monkeypatch.setattr(factory, "get_ocr_extractor", lambda: _StubPdfExtractor())

    res = service._ocr_via_extractor(b"fake-pdf", "scan.pdf", ".pdf", 0.0)

    assert res["pages"] == 3, f"pages 应为总页数 3，实际 {res['pages']}"
    texts = [ln["text"] for ln in res["lines"]]
    assert len(texts) == 6, f"应累积 3 页 × 2 行 = 6 行，实际 {len(texts)}: {texts}"
    assert [ln["page"] for ln in res["lines"]] == [1, 1, 2, 2, 3, 3], "每行都应标注所属页码"
    assert "page-1-A" in texts and "page-3-B" in texts, "首末页文字都应保留"


# ---- 回归测试：PDF 渲染不依赖 PyMuPDF ----

def test_factory_renders_pdf_without_pymupdf():
    """factory 的 PDF 渲染/文本提取必须基于 poppler，而非 PyMuPDF。

    mips64el 等架构没有 PyMuPDF 预编译 wheel（部署时被刻意排除），
    因此本地 OCR 回退路径不能依赖 fitz。
    """
    import inspect

    assert "fitz" not in inspect.getsource(factory), "factory 不应依赖 PyMuPDF(fitz)"

    # poppler 可用时能真实渲染出页面图片
    import shutil as _shutil
    if not _shutil.which("pdftoppm"):
        pytest.skip("环境未安装 pdftoppm（poppler-utils）")

    pdf = (b"%PDF-1.4\n"
           b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
           b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
           b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
           b"trailer<</Root 1 0 R>>\n")
    extractor = factory.LocalPurePythonExtractor()
    pages = extractor.render_pdf_pages(pdf, dpi=72)
    assert len(pages) == 1, f"应渲染出 1 页，实际 {len(pages)}"
    page_no, png_bytes = pages[0]
    assert page_no == 1
    assert png_bytes[:8] == b"\x89PNG\r\n\x1a\n", "渲染结果应为 PNG"